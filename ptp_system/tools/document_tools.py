"""
Outils de génération de documents Word (.docx) et Excel (.xlsx).
Utilisés par le DocumentGeneratorAgent.
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from ptp_system.config import CANDIDAT, FORMATIONS_CIBLES, ETAPES_PTP, INFO_PTP


OUTPUT_DIR = os.getenv("OUTPUT_DIR", "outputs/dossier_ptp")


def _ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


# ---------------------------------------------------------------------------
# Helpers Word
# ---------------------------------------------------------------------------

def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)


def _add_paragraph(doc: Document, text: str, bold: bool = False,
                   indent: bool = False, space_after: int = 6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(1)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def _set_margins(doc: Document, top=2.0, bottom=2.0, left=2.5, right=2.5):
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)


# ---------------------------------------------------------------------------
# 1. Lettre de motivation Transitions Pro
# ---------------------------------------------------------------------------

def generer_lettre_transitions_pro(output_path: str = None) -> str:
    """
    Génère la lettre de motivation pour le dossier Transitions Pro (CPIR).
    Retourne le chemin du fichier créé.
    """
    _ensure_output_dir()
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, "lettre_motivation_transitions_pro.docx")

    c = CANDIDAT
    doc = Document()
    _set_margins(doc)

    # --- En-tête expéditeur ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{c['nom_complet']}\n")
    r.bold = True
    r.font.size = Pt(12)
    p.add_run("Technicien ascensoriste | 4 ans d'expérience\n")
    p.add_run(f"Formation : BTS Électrotechnique + Bac STI2D\n")

    doc.add_paragraph()

    # --- Destinataire ---
    p = doc.add_paragraph()
    p.add_run("À l'attention de\n").bold = True
    p.add_run("Transitions Pro (CPIR)\n")
    p.add_run("Commission Paritaire Interprofessionnelle Régionale")

    doc.add_paragraph()

    # --- Lieu et date ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"Le {datetime.now().strftime('%d %B %Y')}")

    doc.add_paragraph()

    # --- Objet ---
    p = doc.add_paragraph()
    r = p.add_run("Objet : Demande de financement PTP — Licence Professionnelle Automatismes "
                  "et Informatique Industrielle")
    r.bold = True
    r.font.size = Pt(11)

    doc.add_paragraph()

    # --- Corps de la lettre ---
    _add_paragraph(doc, "Madame, Monsieur,", bold=True)

    _add_paragraph(doc,
        f"Technicien ascensoriste depuis {c['experience_annees']} ans, je me permets de vous "
        "adresser cette demande de financement dans le cadre du Projet de Transition "
        "Professionnelle (PTP), afin de réaliser une Licence Professionnelle en "
        "Automatismes et Informatique Industrielle à la rentrée de septembre 2026."
    )

    _add_paragraph(doc,
        "Titulaire d'un BTS Électrotechnique et d'un Bac STI2D, j'ai acquis au fil de mes "
        "missions une solide expertise en électrotechnique, maintenance de systèmes complexes "
        "et diagnostic de pannes. Chaque jour, j'interviens sur des équipements de pointe — "
        "variateurs de fréquence, armoires de commande, réseaux de capteurs — ce qui m'a "
        "permis de développer une lecture technique fine des systèmes automatisés. "
        "Cependant, l'évolution rapide des technologies (IoT industriel, maintenance "
        "prédictive, cybersécurité OT) me conforte dans la nécessité d'une montée en "
        "compétences formalisée pour rester un acteur compétitif sur le marché de l'emploi."
    )

    _add_paragraph(doc,
        "Mon projet professionnel est clairement défini : accéder à des fonctions "
        "d'ingénieur ou technicien expert en automatismes, IoT ou cybersécurité industrielle. "
        "Ces métiers connaissent une tension de recrutement extrêmement forte (45 000 à "
        "65 000 € annuels) et s'inscrivent pleinement dans les enjeux de la transition "
        "numérique et énergétique des bâtiments. La cohérence entre mon parcours terrain "
        "et les formations visées est un atout que je compte valoriser pleinement."
    )

    _add_paragraph(doc,
        "Cette formation d'un an, dispensée en présentiel à temps plein, représente "
        "un investissement stratégique pour ma carrière. Grâce au PTP, je pourrai me "
        "former sans compromettre ma situation financière, tout en apportant à terme une "
        "valeur ajoutée significative à mes futurs employeurs et à l'ensemble du secteur "
        "de la maintenance industrielle connectée."
    )

    _add_paragraph(doc,
        "Je reste bien entendu à votre disposition pour tout entretien ou document "
        "complémentaire que vous jugerez nécessaire à l'instruction de mon dossier."
    )

    doc.add_paragraph()
    _add_paragraph(doc, "Dans l'attente de votre réponse favorable, je vous adresse, "
                   "Madame, Monsieur, l'expression de mes sincères salutations.")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(c["nom_complet"]).bold = True

    doc.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# 2. Lettre de motivation IUT
# ---------------------------------------------------------------------------

def generer_lettre_iut(formation_nom: str = None, etablissement_nom: str = None,
                        output_path: str = None) -> str:
    """
    Génère la lettre de motivation pour candidater dans les IUT.
    Retourne le chemin du fichier créé.
    """
    _ensure_output_dir()
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, "lettre_motivation_iut.docx")

    formation_nom = formation_nom or "Licence Professionnelle Automatismes et Informatique Industrielle"
    etablissement_nom = etablissement_nom or "votre établissement"

    c = CANDIDAT
    doc = Document()
    _set_margins(doc)

    # --- En-tête ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{c['nom_complet']}\n")
    r.bold = True
    r.font.size = Pt(12)
    p.add_run("Technicien ascensoriste | BTS Électrotechnique\n")

    doc.add_paragraph()

    # --- Destinataire ---
    p = doc.add_paragraph()
    p.add_run("À l'attention du Responsable des admissions\n").bold = True
    p.add_run(f"{etablissement_nom}\n")
    p.add_run(f"Département GEII / {formation_nom[:40]}")

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"Le {datetime.now().strftime('%d %B %Y')}")

    doc.add_paragraph()

    # --- Objet ---
    p = doc.add_paragraph()
    r = p.add_run(f"Objet : Candidature — {formation_nom} — Rentrée Septembre 2026")
    r.bold = True

    doc.add_paragraph()

    _add_paragraph(doc, "Madame, Monsieur,", bold=True)

    _add_paragraph(doc,
        f"Technicien ascensoriste depuis {c['experience_annees']} ans au sein d'une société "
        "de maintenance et d'installation d'ascenseurs, je vous adresse ma candidature pour "
        f"intégrer la {formation_nom} dispensée à {etablissement_nom} à la rentrée de "
        "septembre 2026, dans le cadre d'un Projet de Transition Professionnelle (PTP)."
    )

    _add_paragraph(doc,
        "Titulaire d'un BTS Électrotechnique et d'un Bac STI2D, mon quotidien professionnel "
        "m'a conduit à maîtriser des technologies qui sont au cœur de votre formation : "
        "variateurs de fréquence, armoires de commande à API, réseaux de capteurs, "
        "systèmes de supervision à distance. Chaque dépannage d'ascenseur est en réalité "
        "une intervention sur un système automatisé pluritechnologique exigeant rigueur, "
        "méthode et capacité d'analyse rapide — autant de qualités que je m'engage à "
        "mettre au service des apprentissages académiques."
    )

    _add_paragraph(doc,
        "Ce qui me distingue d'un candidat en formation initiale, c'est une maturité "
        "professionnelle forgée sur le terrain : gestion des priorités en situation d'urgence, "
        "dialogue technique avec les équipes de bureau d'études, respect strict des normes "
        "de sécurité NF EN 81, et sens des responsabilités vis-à-vis des utilisateurs "
        "des équipements. Cette expérience sera un atout complémentaire dans les travaux "
        "pratiques et projets techniques de votre formation."
    )

    _add_paragraph(doc,
        "Mon objectif à l'issue de cette formation est de contribuer au développement "
        "des systèmes automatisés connectés, en particulier dans les secteurs de la "
        "maintenance prédictive, de l'IoT industriel ou de la cybersécurité OT — "
        "domaines en forte croissance où l'alliance de compétences terrain et de "
        "connaissances théoriques approfondies représente une réelle valeur ajoutée."
    )

    _add_paragraph(doc,
        "Je serais honoré(e) de pouvoir vous présenter mon projet lors d'un entretien "
        "et reste à votre entière disposition pour tout renseignement complémentaire."
    )

    doc.add_paragraph()
    _add_paragraph(doc, "Veuillez agréer, Madame, Monsieur, l'expression de mes sincères salutations.")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(c["nom_complet"]).bold = True

    doc.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# 3. Planning démarches PTP (.xlsx)
# ---------------------------------------------------------------------------

def generer_planning_ptp(output_path: str = None) -> str:
    """
    Génère le tableau de suivi des démarches PTP en Excel.
    Retourne le chemin du fichier créé.
    """
    _ensure_output_dir()
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, "planning_demarches_ptp.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Planning PTP"

    # --- Titre ---
    ws.merge_cells("A1:F1")
    titre = ws["A1"]
    titre.value = f"Planning Démarches PTP — {CANDIDAT['nom_complet']} — {INFO_PTP['rentree_cible']}"
    titre.font = Font(bold=True, size=14, color="FFFFFF")
    titre.fill = PatternFill("solid", fgColor="1F497D")
    titre.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # --- En-têtes colonnes ---
    headers = ["N°", "Action / Étape", "Organisme", "Délai", "Statut", "Notes"]
    col_widths = [5, 42, 28, 20, 15, 45]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[2].height = 22

    # --- Données ---
    status_colors = {
        "✅ Fait": "C6EFCE",
        "🔄 En cours": "FFEB9C",
        "À faire": "FCE4D6",
    }

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for row_idx, etape in enumerate(ETAPES_PTP, start=3):
        row_data = [
            etape["num"],
            etape["action"],
            etape["organisme"],
            etape["delai"],
            etape["statut"],
            etape["notes"],
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True)

        # Colorier la ligne selon le statut
        statut = etape["statut"]
        for key, color in status_colors.items():
            if key in statut:
                for col_idx in range(1, 7):
                    ws.cell(row=row_idx, column=col_idx).fill = PatternFill("solid", fgColor=color)
                break

        ws.row_dimensions[row_idx].height = 35

    # --- Figer les volets ---
    ws.freeze_panes = "A3"

    # --- Filtres ---
    ws.auto_filter.ref = f"A2:F{len(ETAPES_PTP) + 2}"

    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# 4. Liste IUT contacts (.xlsx)
# ---------------------------------------------------------------------------

def generer_liste_iut_contacts(output_path: str = None) -> str:
    """
    Génère le tableau des établissements à contacter en Excel.
    Retourne le chemin du fichier créé.
    """
    _ensure_output_dir()
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, "liste_iut_contacts.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts IUT"

    # --- Titre ---
    ws.merge_cells("A1:H1")
    titre = ws["A1"]
    titre.value = f"Établissements à contacter — Dossier PTP {CANDIDAT['nom_complet']}"
    titre.font = Font(bold=True, size=13, color="FFFFFF")
    titre.fill = PatternFill("solid", fgColor="1F497D")
    titre.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # --- En-têtes ---
    headers = [
        "Établissement", "Ville", "Région",
        "Formation", "Email contact",
        "Lien candidature", "Date limite", "Statut"
    ]
    col_widths = [22, 15, 22, 38, 32, 38, 15, 15]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF", size=10)

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[2].height = 22

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # Alternance de couleurs par formation (priorité)
    formation_colors = ["EBF3FB", "FFF2CC", "E2EFDA", "FCE4D6", "F2F2F2"]

    row_idx = 3
    for form_idx, formation in enumerate(FORMATIONS_CIBLES):
        color = formation_colors[form_idx % len(formation_colors)]
        for etab in formation["etablissements"]:
            row_data = [
                etab["nom"],
                etab["ville"],
                etab["region"],
                formation["nom"],
                etab.get("email_contact", "À rechercher"),
                etab.get("lien_candidature", etab["url"]),
                etab.get("date_limite", "À confirmer"),
                "À contacter",
            ]
            for col_idx_data, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx_data, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor=color)

            ws.row_dimensions[row_idx].height = 28
            row_idx += 1

    # --- Mise en forme conditionnelle colonne Statut ---
    last_row = row_idx - 1
    status_col = "H"
    ws.conditional_formatting.add(
        f"{status_col}3:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"✅ Contacté"'],
                   fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    ws.conditional_formatting.add(
        f"{status_col}3:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"🔄 En attente"'],
                   fill=PatternFill("solid", fgColor="FFEB9C")),
    )
    ws.conditional_formatting.add(
        f"{status_col}3:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"❌ Refusé"'],
                   fill=PatternFill("solid", fgColor="FFC7CE")),
    )

    # --- Figer et filtres ---
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:H{last_row}"

    wb.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Génération de tous les documents
# ---------------------------------------------------------------------------

def generer_tous_documents() -> dict[str, str]:
    """
    Génère l'ensemble des documents PTP.
    Retourne un dict {nom_doc: chemin_fichier}.
    """
    resultats = {}

    resultats["lettre_transitions_pro"] = generer_lettre_transitions_pro()
    resultats["lettre_iut"] = generer_lettre_iut()
    resultats["planning_ptp"] = generer_planning_ptp()
    resultats["liste_contacts"] = generer_liste_iut_contacts()

    return resultats
